"""Reproduce official attachments; no tender or company data belongs in this package."""
import json, re, hashlib, argparse
from pathlib import Path
import pdfplumber
from docx import Document
BASE=Path(__file__).parent
parser=argparse.ArgumentParser(description='Rebuild the catalog from explicitly cached, hash-verified official sources.')
parser.add_argument('--source-dir',type=Path,default=BASE.parent.parent/'policy-source-cache')
parser.add_argument('--output',type=Path,default=BASE/'cn-gpp-2019-2025.1.json')
args=parser.parse_args()
OFF=args.source_dir
EXPECTED={s['id']:s['sha256'] for s in json.loads((BASE/'cn-gpp-2019-2025.1.json').read_text())['sources']}
checked='2026-10-14'
sources=[]
def source(id,title,url,file,published,effective,pages=None):
 value=dict(id=id,title=title,url=url,publishedAt=published,effectiveFrom=effective,checkedAt=checked,sha256=hashlib.sha256((OFF/file).read_bytes()).hexdigest(),localFile='official/'+file)
 if pages:value['pageCount']=pages
 if value['sha256']!=EXPECTED[id]:raise ValueError('Pinned official source hash mismatch: '+id)
 sources.append(value)
source('authority-2019','财库〔2019〕9号：节能、环境标志采购执行机制','https://gks.mof.gov.cn/ztztz/zhengfucaigouguanli/201902/t20190212_3146225.htm','authority-2019.html','2019-02-12','2019-04-01')
source('energy-2019','节能产品政府采购品目清单（财库〔2019〕19号附件）','https://gks.mof.gov.cn/ztztz/zhengfucaigouguanli/201904/P020190404409422168443.pdf','energy-2019.pdf','2019-04-03','2019-04-01',4)
sources[-1]['issuedAt']='2019-04-02'
source('environment-2019','环境标志产品政府采购品目清单（财库〔2019〕18号附件）','https://gks.mof.gov.cn/ztztz/zhengfucaigouguanli/201903/P020190329702478575295.pdf','environment-2019.pdf','2019-03-30','2019-04-01',4)
sources[-1]['issuedAt']='2019-03-29'
source('green-authority-2025','财库〔2024〕36号：扩大绿色建材政策实施范围','https://gks.mof.gov.cn/guizhangzhidu/202501/t20250110_3951449.htm','green-material-2025.html','2025-01-10','2025-01-01')
source('green-standard-2025','绿色建筑和绿色建材政府采购需求标准（2025年版）','https://gks.mof.gov.cn/guizhangzhidu/202501/P020250110319997903210.docx','green-material-standard-2025.docx','2025-01-10','2025-01-01')
source('green-cities-2025','政府采购支持绿色建材政策实施城市名单（101市、区）','https://gks.mof.gov.cn/guizhangzhidu/202501/P020250110319997676870.docx','green-material-cities-2025.docx','2025-01-10','2025-01-01')
catalog=[]
def entry(kind,code,name,standard,page,mandatory=False,scope='',ordinal=None,aliases=None,exclude=None,**more):
 e=dict(id=f'{kind}-{len([x for x in catalog if x["policyType"]==kind])+1:03}',categoryCode=code,name=name,policyType=kind,policyKind='requirement' if mandatory else 'preference',standard=standard,scope=scope or name,sourceId={'energy':'energy-2019','environment':'environment-2019','greenMaterial':'green-standard-2025'}[kind],sourcePage=page,aliases=aliases or [],excludedTerms=exclude or [])
 if ordinal:e['ordinal']=ordinal
 e.update(more);catalog.append(e)
# Each row below corresponds to a leaf and its inherited star on the original PDF.
E=[
(1,1,'A02010104','台式计算机','《微型计算机能效限定值及能效等级》（GB 28380）',1),
(1,1,'A02010105','便携式计算机','《微型计算机能效限定值及能效等级》（GB 28380）',1),
(1,1,'A02010107','平板式微型计算机','《微型计算机能效限定值及能效等级》（GB 28380）',1),
(1,2,'A0201060101','喷墨打印机','《复印机、打印机和传真机能效限定值及能效等级》（GB 21521）',0),
(1,2,'A0201060102','激光打印机','《复印机、打印机和传真机能效限定值及能效等级》（GB 21521）',1),
(1,2,'A0201060104','针式打印机','《复印机、打印机和传真机能效限定值及能效等级》（GB 21521）',1),
(1,2,'A0201060401','液晶显示器','《计算机显示器能效限定值及能效等级》（GB 21520）',1),
(1,2,'A0201060901','扫描仪','参照《复印机、打印机和传真机能效限定值及能效等级》（GB 21521）中打印速度为15页/分的针式打印机相关要求',0),
(1,3,'A020202','投影仪','《投影机能效限定值及能效等级》（GB 32028）',0),
(1,4,'A020204','多功能一体机','《复印机、打印机和传真机能效限定值及能效等级》（GB 21521）',0),
(1,5,'A02051901','离心泵','《清水离心泵能效限定值及节能评价值》（GB 19762）',0),
(1,6,'A02052301','冷水机组','《冷水机组能效限定值及能效等级》（GB 19577），《低环境温度空气源热泵（冷水）机组能效限定值及能效等级》（GB 37480）',1),
(1,6,'A02052301','水源热泵机组','《水（地）源热泵机组能效限定值及能效等级》（GB 30721）',1),
(2,6,'A02052301','溴化锂吸收式冷水机组','《溴化锂吸收式冷水机组能效限定值及能效等级》（GB 29540）',1),
(2,6,'A02052305','多联式空调（热泵）机组（制冷量>14000W）','《多联式空调（热泵）机组能效限定值及能源效率等级》（GB 21454）',1),
(2,6,'A02052305','单元式空气调节机（制冷量>14000W）','《单元式空气调节机能效限定值及能效等级》（GB 19576）《风管送风式空调机组能效限定值及能效等级》（GB 37479）',1),
(2,6,'A02052309','机房空调','《单元式空气调节机能效限定值及能效等级》（GB 19576）',1),
(2,6,'A02052399','冷却塔','《机械通风冷却塔 第1部分：中小型开式冷却塔》（GB/T 7190.1）；《机械通风冷却塔 第2部分：大型开式冷却塔》（GB/T 7190.2）',0),
(2,7,'A020601','电机','《中小型三相异步电动机能效限定值及能效等级》（GB 18613）',0),
(2,8,'A020602','配电变压器','《三相配电变压器能效限定值及能效等级》（GB 20052）',0),
(2,9,'A020609','管型荧光灯镇流器','《管形荧光灯镇流器能效限定值及能效等级》（GB 17896）',1),
(2,10,'A0206180101','电冰箱','《家用电冰箱耗电量限定值及能效等级》（GB 12021.2）',0),
(2,10,'A0206180203','房间空气调节器','《转速可控型房间空气调节器能效限定值及能效等级》（GB 21455-2013），待2019年修订发布后，按《房间空气调节器能效限定值及能效等级》（GB21455-2019）实施。',1),
(2,10,'A0206180203','多联式空调（热泵）机组（制冷量≤14000W）','《多联式空调（热泵）机组能效限定值及能源效率等级》（GB 21454）',1),
(2,10,'A0206180203','单元式空气调节机（制冷量≤14000W）','《单元式空气调节机能效限定值及能源效率等级》（GB 19576）《风管送风式空调机组能效限定值及能效等级》（GB 37479）',1),
(2,10,'A0206180301','洗衣机','《电动洗衣机能效水效限定值及等级》（GB 12021.4）',0),
(3,10,'A02061808','电热水器','《储水式电热水器能效限定值及能效等级》（GB 21519）',1),
(3,10,'A02061808','燃气热水器','《家用燃气快速热水器和燃气采暖热水炉能效限定值及能效等级》（GB 20665）',0),
(3,10,'A02061808','热泵热水器','《热泵热水机（器）能效限定值及能效等级》（GB 29541）',0),
(3,10,'A02061808','太阳能热水系统','《家用太阳能热水系统能效限定值及能效等级》（GB 26969）',0),
(3,11,'A020619','普通照明用双端荧光灯','《普通照明用双端荧光灯能效限定值及能效等级》（GB 19043）',1),
(3,11,'A020619','LED道路/隧道照明产品','《道路和隧道照明用LED灯具能效限定值及能效等级》（GB 37478）',0),
(3,11,'A020619','LED筒灯','《室内照明用LED产品能效限定值及能效等级》（GB 30255）',0),
(3,11,'A020619','普通照明用非定向自镇流LED灯','《室内照明用LED产品能效限定值及能效等级》（GB 30255）',0),
(3,12,'A02091001','普通电视设备（电视机）','《平板电视能效限定值及能效等级》（GB 24850）',1),
(3,13,'A02091107','监视器','以射频信号为主要信号输入的监视器应符合《平板电视能效限定值及能效等级》（GB 24850），以数字信号为主要信号输入的监视器应符合《计算机显示器能效限定值及能效等级》（GB 21520）',1),
(3,14,'A031210','商用燃气灶具','《商用燃气灶具能效限定值及能效等级》（GB 30531）',0),
(3,15,'A060805','坐便器','《坐便器水效限定值及水效等级》（GB 25502）',1),
(3,15,'A060805','蹲便器','《蹲便器用水效率限定值及用水效率等级》（GB 30717）',1),
(3,15,'A060805','小便器','《小便器用水效率限定值及用水效率等级》（GB 28377）',1),
(4,16,'A060806','水嘴','《水嘴用水效率限定值及用水效率等级》（GB 25501）',1),
(4,17,'A060807','便器冲洗阀','《便器冲洗阀用水效率限定值及用水效率等级》（GB 28379）',0),
(4,18,'A060810','淋浴器','《淋浴器用水效率限定值及用水效率等级》（GB 28378）',0)]
aliases={'台式计算机':['台式电脑'],'便携式计算机':['笔记本电脑'],'平板式微型计算机':['平板电脑'],'水嘴':['水龙头'],'坐便器':['座便器','马桶'],'蹲便器':['蹲厕'],'电冰箱':['冰箱'],'房间空气调节器':['分体空调','壁挂式空调'],'普通电视设备（电视机）':['电视机'],'水源热泵机组':['地源热泵机组'],'普通照明用双端荧光灯':['双端荧光灯'],'管型荧光灯镇流器':['管形荧光灯镇流器'],'LED道路/隧道照明产品':['LED道路灯','LED隧道灯'],'普通照明用非定向自镇流LED灯':['非定向自镇流LED灯']}
for page,ordinal,code,name,standard,mandatory in E:
 scope={'电机':'中小型三相异步电动机','房间空气调节器':'房间空气调节器；不得仅凭空调机总称判定具体范围','监视器':'射频/数字信号输入按相应标准核对'}.get(name,name)
 entry('energy',code,name,standard,page,bool(mandatory),scope,ordinal,aliases.get(name,[]),['雨水口','雨水斗'] if name=='水嘴' else [])
# Environment PDF has ordinary leaf rows; choose rightmost category cell.
with pdfplumber.open(OFF/'environment-2019.pdf') as pdf:
 for page in pdf.pages:
  ordinal=None
  for row in page.extract_tables()[0]:
   if row[0] and row[0].strip().isdigit():ordinal=int(row[0])
   if not ordinal:continue
   raw=next((s for s in reversed(row[1:4]) if s and re.search(r'A\d+',s)),None)
   if not raw:continue
   raw=raw.replace('\n','').strip();m=re.match(r'(A\d+)\s*(.*)',raw)
   name=m[2];standard=row[4].replace('\n','')
   scope={'A02061808':'仅太阳能集热器（HJ/T362），热水器总称不能证明范围','A020301':'仅HJ2532轻型汽车适用范围','A02030501':'仅HJ2532轻型汽车适用范围','A02030599':'仅HJ2532轻型汽车适用范围','A02030601':'仅HJ2532轻型汽车适用范围','A02030799':'仅HJ2532轻型汽车适用范围','A060804':'仅HJ/T296卫生陶瓷适用范围','A10039901':'仅HJ456刚性防水材料适用范围','A180201':'仅HJ/T226建筑用塑料管材或HJ/T231再生塑料制品','A170112':'仅HJ2541胶粘剂适用范围'}.get(m[1],name+'；认证应依据相关标准最新版本')
   aa=aliases.get(name,[])
   if m[1]=='A10030901':aa=['改性沥青防水卷材','高聚物改性沥青防水卷','SBS防水卷材','SBS改性沥青防水卷材']
   if m[1]=='A10030906':aa=['高分子防水卷材']
   if m[1]=='A10030301':aa=['预拌混凝土']
   if m[1]=='A10060202':aa=['乳胶漆','内墙乳胶漆']
   if m[1]=='A10060499':aa=['水性防水涂料','水性橡胶沥青防水涂料','水性橡胶沥青涂料']
   if m[1]=='A180201':aa=['塑料管材','PVC管','PPR管','PE管','再生塑料制品']
   if m[1]=='A10031001':aa=['岩棉','玻璃棉']
   entry('environment',m[1],name,standard,page.page_number,scope=scope,ordinal=ordinal,aliases=aa,exclude=['雨水口','雨水斗'] if name=='水嘴' else [])
# Official table 1 is a leaf list. Bold denotes mandatory, clause numbers are
# document references, NOT claimed as national procurement category codes.
doc=Document(OFF/'green-material-standard-2025.docx')
clauses={}
for p in doc.paragraphs:
 m=re.match(r'([45]\.\d+\.\d+)\s+(.+?)\t(\d+)$',p.text)
 if m:clauses[m[2].strip()]=(m[1],int(m[3]))
for rownum,row in enumerate(doc.tables[0].rows[1:-1],1):
 middle=row.cells[2].text.removesuffix('a');clause,page=clauses[middle]
 names=row.cells[3].text.removesuffix('a').split('、')
 mandatory=clause.startswith('4.')
 for leaf,name in enumerate(names,1):
  more={'sourceTable':1,'sourceRow':rownum+1,'clause':clause,'group':middle}
  if middle=='钢结构构件':more['conditions']={'buildingStructure':'钢结构'}
  if middle=='混凝土结构构件':more['conditions']={'buildingStructure':'装配式混凝土结构'}
  # Preserve the entire original standard; detailed numeric requirements are
  # not reduced to invented universal pass rules.
  entry('greenMaterial',f'GM2025:{clause}:{leaf}',name,'《绿色建筑和绿色建材政府采购需求标准（2025年版）》第'+clause+'条',page,mandatory,scope='按该小类全部适用技术指标和地方标准核对；具体指标见来源链接中的官方DOCX（本源码不附原件）',aliases={'门窗':['金属窗','铝合金窗','塑料窗','天窗','断桥铝窗','铝合金门窗'],'岩棉制品':['岩棉板'],'改性沥青防水卷材':['高聚物改性沥青防水卷','SBS防水卷材','SBS改性沥青防水卷材'],'水性防水涂料':['水性橡胶沥青防水涂料','水性橡胶沥青涂料'],'挤塑聚苯乙烯泡沫塑料制品（XPS）':['XPS','挤塑聚苯板'],'模塑聚苯乙烯泡沫塑料制品（EPS）':['EPS','模塑聚苯板'],'水性墙面涂料':['内墙乳胶漆'],'水泥基预拌砂浆':['预拌水泥砂浆']}.get(name,[]),exclude=['雨水口','雨水斗'] if name=='水嘴' else [],**more)
cities=[];province=None
for p in Document(OFF/'green-material-cities-2025.docx').paragraphs:
 t=p.text.strip()
 if re.match(r'^[一二三四五六七八九十]+、',t):province=t.split('、',1)[1]
 m=re.match(r'^(\d+)\.(.+)$',t)
 if m:cities.append(dict(province=province,name=m[2],ordinal=int(m[1])))
pkg={'schemaVersion':1,'version':'cn-gpp-2019-2025.1-oss.1','checkedAt':checked,'sources':sources,'catalog':catalog,'notes':['节能、环境标志编码保留2019官方附件原编码，不推断为现行采购分类编码。','2019执行机制4月1日生效，节能目录19号4月2日发文、4月3日发布，历史时点不得混同。','标准列为官方目录原文；适用证书应依据相关标准最新版本；离线包核对日期之后的变更须显式核查。','绿色GM2025编码为包内目录引用编号，sourcePage为官方需求标准印刷页码（非DOCX渲染页）；具体指标以来源链接中的官方DOCX为准，本源码不附原件。'],'greenMaterial':{'sourceId':'green-standard-2025','citiesSourceId':'green-cities-2025','authoritySourceId':'green-authority-2025','effectiveFrom':'2025-01-01','cities':cities,'projectTypes':['医院','学校','办公楼','综合体','展览馆','会展中心','体育馆','保障性住房','旧城改造'],'optionalCategoryRatio':0.4,'withinCategoryRatio':0.8}}
args.output.write_text(json.dumps(pkg,ensure_ascii=False,indent=2)+'\n')
print({k:len([x for x in catalog if x['policyType']==k]) for k in ['energy','environment','greenMaterial']},'cities',len(cities))
